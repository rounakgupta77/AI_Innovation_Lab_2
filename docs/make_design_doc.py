"""Builds docs/Technical_Design_Document.docx from docs/Technical_Design_Document.md.

The four flow diagrams are drawn here with Pillow (saved in docs/_diagrams/ as PNG), and the picture of the
block-diagram slide (docs/preview.png, made by make_block_diagram.py --preview) is added if it exists.

    python docs/make_design_doc.py --track "Your track name"
"""
import argparse
import math
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
MD_PATH = os.path.join(HERE, "Technical_Design_Document.md")
OUT_PATH = os.path.join(HERE, "Technical_Design_Document.docx")
DIAGRAM_DIR = os.path.join(HERE, "_diagrams")
PREVIEW = os.path.join(HERE, "preview.png")
PLACEHOLDER = "`[TRACK NAME: replace this line with your track]`"

INK, ARROW_C, GREY = (29, 36, 51), (74, 85, 104), (107, 116, 134)
BLUE, BLUE_T = (11, 92, 173), (233, 241, 250)
GREEN, GREEN_T = (18, 122, 70), (232, 245, 238)
PURPLE, PURPLE_T = (91, 60, 196), (239, 235, 250)
RED, RED_T = (180, 35, 24), (253, 236, 234)
AMBER, AMBER_T = (154, 91, 0), (255, 243, 214)
GREY_T, WHITE = (238, 240, 244), (255, 255, 255)

TEXT_WIDTH_IN = 6.85


# --------------------------------------------------------------------------- #
# Diagram drawing (Pillow)
# --------------------------------------------------------------------------- #

def font(size, bold=False):
    names = ["calibrib.ttf", "arialbd.ttf", "DejaVuSans-Bold.ttf"] if bold else ["calibri.ttf", "arial.ttf", "DejaVuSans.ttf"]
    for name in names:
        for folder in (r"C:\Windows\Fonts", ""):
            try:
                return ImageFont.truetype(os.path.join(folder, name) if folder else name, size)
            except OSError:
                continue
    return ImageFont.load_default()


def new_canvas(w, h):
    img = Image.new("RGB", (w, h), WHITE)
    return img, ImageDraw.Draw(img)


def text_center(d, cx, cy, text, size=24, bold=False, color=INK):
    f, lines, lh = font(size, bold), text.split("\n"), int(size * 1.28)
    top = cy - lh * len(lines) / 2 + lh / 2
    for i, line in enumerate(lines):
        d.text((cx, top + i * lh), line, font=f, fill=color, anchor="mm")


def box(d, x, y, w, h, text, fill, line, size=24, bold=False, color=INK, kind="rr"):
    if kind == "diamond":
        pts = [(x + w / 2, y), (x + w, y + h / 2), (x + w / 2, y + h), (x, y + h / 2)]
        d.polygon(pts, fill=fill)
        d.line(pts + [pts[0]], fill=line, width=3)
    else:
        d.rounded_rectangle([x, y, x + w, y + h], radius=18, fill=fill, outline=line, width=3)
    text_center(d, x + w / 2, y + h / 2, text, size, bold, color)


def label(d, cx, cy, text, size=21, color=GREY, bold=False):
    f, lines, lh = font(size, bold), text.split("\n"), int(size * 1.25)
    width = max(d.textlength(line, font=f) for line in lines)
    d.rectangle([cx - width / 2 - 5, cy - lh * len(lines) / 2 - 2, cx + width / 2 + 5, cy + lh * len(lines) / 2 + 2], fill=WHITE)
    text_center(d, cx, cy, text, size, bold, color)


def _head(d, frm, to, color, size=17):
    ang = math.atan2(to[1] - frm[1], to[0] - frm[0])
    a = (to[0] - size * math.cos(ang - 0.4), to[1] - size * math.sin(ang - 0.4))
    b = (to[0] - size * math.cos(ang + 0.4), to[1] - size * math.sin(ang + 0.4))
    d.polygon([to, a, b], fill=color)


def arrow(d, pts, both=False, dashed=False, color=ARROW_C):
    for p, q in zip(pts, pts[1:]):
        length = math.dist(p, q)
        if not dashed:
            d.line([p, q], fill=color, width=3)
            continue
        step, on = 24, 14
        for s in range(0, int(length), step):
            e = min(s + on, length)
            d.line([(p[0] + (q[0] - p[0]) * s / length, p[1] + (q[1] - p[1]) * s / length),
                    (p[0] + (q[0] - p[0]) * e / length, p[1] + (q[1] - p[1]) * e / length)], fill=color, width=3)
    _head(d, pts[-2], pts[-1], color)
    if both:
        _head(d, pts[1], pts[0], color)


def diagram_architecture():
    img, d = new_canvas(1500, 800)
    d.rounded_rectangle([400, 30, 1190, 770], radius=24, fill=GREEN_T, outline=GREEN, width=3)
    d.text((430, 64), "SHARED BACKEND  (Flask, port 5000)", font=font(25, True), fill=GREEN, anchor="lm")
    box(d, 40, 120, 230, 80, "Employee\nbrowser", WHITE, GREY)
    box(d, 40, 260, 230, 110, "Employee portal\nHTML + JS, port 5001", BLUE_T, BLUE)
    box(d, 40, 430, 230, 110, "HR portal\nHTML + JS, port 5002", PURPLE_T, PURPLE)
    box(d, 40, 600, 230, 80, "HR\nbrowser", WHITE, GREY)
    arrow(d, [(155, 200), (155, 260)])
    arrow(d, [(155, 600), (155, 540)])
    box(d, 440, 300, 300, 160, "Flask API\nlogin, bills,\nHR decisions", WHITE, GREEN, bold=True)
    box(d, 440, 100, 300, 100, "Uploaded bill files\n(images and PDFs)", GREY_T, GREY)
    box(d, 820, 100, 320, 130, "Bill reader\nClaude AI if a key is set,\nelse Tesseract OCR", WHITE, GREEN)
    box(d, 820, 330, 320, 130, "Anomaly detection\nstatistical rules +\nIsolation Forest", RED_T, RED)
    box(d, 820, 600, 320, 100, "SQLite database\nbills, users, settings", AMBER_T, AMBER)
    box(d, 1250, 100, 220, 130, "Anthropic API\n(optional,\nneeds internet)", WHITE, GREY)
    arrow(d, [(270, 315), (440, 350)], both=True)
    arrow(d, [(270, 485), (440, 410)], both=True)
    label(d, 355, 382, "REST / JSON")
    arrow(d, [(590, 300), (590, 200)])
    arrow(d, [(740, 150), (820, 150)])
    arrow(d, [(740, 320), (860, 230)])
    arrow(d, [(740, 380), (820, 380)])
    arrow(d, [(980, 460), (980, 600)], both=True)
    arrow(d, [(590, 460), (590, 650), (820, 650)])
    arrow(d, [(1140, 165), (1250, 165)], both=True, dashed=True)
    text_center(d, 1360, 330, "Dashed: optional.\nSolid: always used.", 21, False, GREY)
    return img


def diagram_flow():
    img, d = new_canvas(1500, 790)
    row1 = [(40, "Employee uploads\nbills", BLUE_T, BLUE), (340, "Backend saves\nthe files (status:\nprocessing)", GREEN_T, GREEN),
            (640, "Bill reader fills in\nnumber, date,\namount, type", GREEN_T, GREEN),
            (940, "Employee reviews\nand corrects fields", BLUE_T, BLUE), (1240, "Employee clicks\nSend to HR", BLUE_T, BLUE)]
    for x, text, fill, line in row1:
        box(d, x, 60, 240, 130, text, fill, line, size=23)
    for x, *_ in row1[:-1]:
        arrow(d, [(x + 240, 125), (x + 300, 125)])
    arrow(d, [(1360, 190), (1360, 285), (160, 285), (160, 400)])
    box(d, 40, 400, 240, 130, "Backend validates,\nruns anomaly check", GREEN_T, GREEN, size=23)
    box(d, 320, 370, 340, 190, "Auto-approve ON\nand not flagged?", AMBER_T, AMBER, size=22, bold=True, kind="diamond")
    box(d, 720, 400, 220, 130, "Waiting for HR\n(red if flagged)", PURPLE_T, PURPLE, size=23)
    box(d, 990, 400, 220, 130, "HR approves\nor rejects", PURPLE_T, PURPLE, size=23)
    box(d, 1260, 400, 210, 130, "Employee sees\nstatus", BLUE_T, BLUE, size=23)
    box(d, 380, 640, 220, 100, "Approved\nautomatically", GREEN_T, GREEN, size=23)
    arrow(d, [(280, 465), (320, 465)])
    arrow(d, [(660, 465), (720, 465)])
    arrow(d, [(940, 465), (990, 465)])
    arrow(d, [(1210, 465), (1260, 465)])
    arrow(d, [(490, 560), (490, 640)])
    arrow(d, [(600, 690), (1365, 690), (1365, 530)])
    label(d, 690, 440, "No", 22, INK, True)
    label(d, 525, 600, "Yes", 22, GREEN, True)
    return img


def diagram_lifecycle():
    img, d = new_canvas(1500, 700)
    d.ellipse([28, 293, 52, 317], fill=INK)
    box(d, 90, 250, 220, 110, "Processing\n(file uploaded)", GREY_T, GREY)
    box(d, 470, 250, 220, 110, "Draft\n(not sent)", BLUE_T, BLUE)
    box(d, 870, 250, 220, 110, "Pending\n(with HR)", AMBER_T, AMBER)
    box(d, 1230, 60, 240, 100, "Approved", GREEN_T, GREEN, bold=True)
    box(d, 1230, 450, 240, 100, "Rejected", RED_T, RED, bold=True)
    box(d, 470, 470, 220, 90, "Deleted", WHITE, GREY)
    arrow(d, [(52, 305), (90, 305)])
    arrow(d, [(310, 305), (470, 305)])
    label(d, 390, 268, "reader\nfinishes")
    arrow(d, [(690, 305), (870, 305)])
    label(d, 780, 268, "employee\nsends to HR")
    arrow(d, [(1090, 290), (1230, 145)])
    label(d, 1200, 232, "HR approves")
    arrow(d, [(1090, 320), (1230, 495)])
    label(d, 1030, 440, "HR rejects\n(reason required)")
    arrow(d, [(580, 250), (580, 115), (1230, 115)])
    label(d, 900, 70, "sent while auto-approve is ON\nand the bill is not flagged")
    arrow(d, [(580, 360), (580, 470)])
    label(d, 690, 415, "employee deletes")
    text_center(d, 750, 640, "Bills already sent to HR can no longer be edited or deleted.", 22, False, GREY)
    return img


def diagram_anomaly():
    img, d = new_canvas(1500, 740)
    box(d, 40, 300, 210, 140, "Bill sent\n(type, date,\namount)", BLUE_T, BLUE)
    box(d, 310, 290, 250, 160, "Load history:\napproved bills +\nunflagged pending\nbills", GREY_T, GREY, size=23)
    checks = ["1. Own history\nvs this employee's bills\nof the same type", "2. Company\nvs everyone's bills\nof the same type",
              "3. Monthly total\nemployee's month total\nfor the type", "4. Isolation Forest\nML outlier model\n(30+ bills of history)"]
    centers = []
    for i, text in enumerate(checks):
        y = 40 + i * 175
        box(d, 620, y, 340, 130, text, GREEN_T if i < 3 else RED_T, GREEN if i < 3 else RED, size=22)
        centers.append(y + 65)
    box(d, 1020, 270, 200, 200, "Any check\nfired?", AMBER_T, AMBER, size=23, bold=True, kind="diamond")
    box(d, 1280, 140, 190, 170, "Flagged:\nred row on\nHR portal\nwith reasons", RED_T, RED, size=22)
    box(d, 1280, 430, 190, 130, "Not flagged", GREEN_T, GREEN, size=22)
    arrow(d, [(250, 370), (310, 370)])
    for cy in centers:
        arrow(d, [(560, 370), (620, cy)])
        arrow(d, [(960, cy), (1030, 370)])
    arrow(d, [(1120, 270), (1120, 225), (1280, 225)])
    arrow(d, [(1120, 470), (1120, 495), (1280, 495)])
    label(d, 1200, 205, "Yes", 22, RED, True)
    label(d, 1200, 478, "No", 22, GREEN, True)
    return img


DIAGRAMS = {"architecture": diagram_architecture, "flow": diagram_flow,
            "lifecycle": diagram_lifecycle, "anomaly": diagram_anomaly}


# --------------------------------------------------------------------------- #
# Word helpers
# --------------------------------------------------------------------------- #

def set_style_font(style, name, size=None, color=None, bold=None):
    style.font.name = name
    style.element.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:eastAsia"), name)
    if size:
        style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def add_inline(p, text, size=None, bold=False, italic=False, color=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        run_bold, run_italic, code = bold, italic, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, run_bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, run_italic = part[1:-1], True
        run = p.add_run(part)
        run.bold, run.italic = run_bold, run_italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if code:
            run.font.name = "Consolas"
            run.element.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt((size or 10.5) - 1)
            run.font.color.rgb = RGBColor.from_string("7A1F2B")


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    shade_paragraph(p, "F2F4F8")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.1)
    for i, line in enumerate(lines):
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.element.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.5)
        if i < len(lines) - 1:
            run.add_break()


def add_table(doc, rows):
    header = rows[0]
    has_header = any(c.strip() for c in header)
    body = rows[1:]
    all_rows = ([header] if has_header else []) + body
    cols = len(header)
    weights = [min(max(max(len(re.sub(r"[*`]", "", r[c])) for r in all_rows if c < len(r)), 8), 55) for c in range(cols)]
    total = sum(weights)
    widths = [Inches(TEXT_WIDTH_IN * w / total) for w in weights]

    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_index, row in enumerate(all_rows):
        cells = table.add_row().cells
        for c in range(cols):
            cell = cells[c]
            cell.width = widths[c]
            if has_header and r_index == 0:
                shade_cell(cell, "DCE6F4")
            elif r_index % 2 == 0:
                shade_cell(cell, "F7F9FC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            add_inline(p, row[c] if c < len(row) else "", size=9.5, bold=has_header and r_index == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_picture(doc, path, width_in=TEXT_WIDTH_IN):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width_in))


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Fuel Reimbursement Management System · Technical Design Document · page ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string("6B7486")
    num = p.add_run()
    num.font.size = Pt(8.5)
    num.font.color.rgb = RGBColor.from_string("6B7486")
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        num._r.append(el)


# --------------------------------------------------------------------------- #
# Markdown -> Word
# --------------------------------------------------------------------------- #

def build(track):
    os.makedirs(DIAGRAM_DIR, exist_ok=True)
    images = {}
    for name, draw in DIAGRAMS.items():
        path = os.path.join(DIAGRAM_DIR, f"{name}.png")
        draw().save(path)
        images[name] = path

    with open(MD_PATH, encoding="utf-8") as fh:
        text = fh.read()
    if track:
        text = text.replace(PLACEHOLDER, track)
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.left_margin = section.right_margin = Inches(0.7)
    section.top_margin, section.bottom_margin = Inches(0.8), Inches(0.8)
    add_page_number_footer(section)

    set_style_font(doc.styles["Normal"], "Calibri", 10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    set_style_font(doc.styles["Title"], "Calibri", 26, "1F3864", True)
    set_style_font(doc.styles["Heading 1"], "Calibri", 17, "1F3864", True)
    set_style_font(doc.styles["Heading 2"], "Calibri", 13.5, "2F5597", True)
    set_style_font(doc.styles["Heading 3"], "Calibri", 11.5, "2F5597", True)

    pending_diagram, i = None, 0
    figure_no = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s or s == "---":
            i += 1
            continue

        m = re.match(r"<!-- diagram: (\w+) -->", s)
        if m:
            pending_diagram = m.group(1)
            i += 1
            continue
        if s == "<!-- slide-preview -->":
            if os.path.exists(PREVIEW):
                add_picture(doc, PREVIEW)
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(cap, "Figure 3: The block diagram slide (Reimbursement_Block_Diagram.pptx, attached with this document)",
                           size=9, italic=True, color="6B7486")
            else:
                note = doc.add_paragraph()
                add_inline(note, "The slide image could not be generated automatically on this computer. "
                                 "See the attached Reimbursement_Block_Diagram.pptx.", italic=True, color="6B7486")
            i += 1
            continue

        if s.startswith("```"):
            lang, block, i = s[3:].strip(), [], i + 1
            while not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            if lang == "mermaid" and pending_diagram in images:
                add_picture(doc, images[pending_diagram])
                figure_no += 1
                pending_diagram = None
            else:
                add_code_block(doc, block)
            continue

        if s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            heading = s[level:].strip()
            if level == 1:
                doc.add_heading(heading, 0)
            else:
                doc.add_heading(heading, level - 1)
            i += 1
            continue

        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        if s.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, lines[i].strip()[2:])
                i += 1
            continue

        if re.match(r"\d+\. ", s):
            while i < len(lines) and re.match(r"\d+\. ", lines[i].strip()):
                number, rest = lines[i].strip().split(". ", 1)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, f"{number}.  {rest}")
                i += 1
            continue

        if s.startswith("> "):
            p = doc.add_paragraph()
            shade_paragraph(p, "FFF3D6")
            p.paragraph_format.left_indent = Inches(0.1)
            add_inline(p, s[2:], italic=False)
            i += 1
            continue

        p = doc.add_paragraph()
        if s.startswith("*Figure"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, s, size=9, color="6B7486")
        else:
            add_inline(p, s)
        i += 1

    doc.core_properties.title = "Technical Design Document: Fuel Reimbursement Management System"
    doc.save(OUT_PATH)
    print("Saved", OUT_PATH)
    print("Diagram images saved in", DIAGRAM_DIR)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--track", default="", help="Name of the track, written into the project overview")
    build(ap.parse_args().track.strip())
