"""Builds docs/Reimbursement_Block_Diagram.pptx: one slide with the system's block diagram.

    python docs/make_block_diagram.py            build the slide
    python docs/make_block_diagram.py --preview  also export docs/preview.png (needs PowerPoint, Windows)
"""
import os
import subprocess
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Reimbursement_Block_Diagram.pptx")
PREVIEW = os.path.join(HERE, "preview.png")

INK, MUTED, ARROW = "1D2433", "5A6578", "4A5568"
BLUE, GREEN, PURPLE, RED, AMBER, GREY = "0B5CAD", "127A46", "5B3CC4", "B42318", "9A5B00", "6B7486"
BLUE_T, GREEN_T, PURPLE_T = "E9F1FA", "E8F5EE", "EFEBFA"
RED_T, AMBER_T, GREY_T, ROW, WHITE = "FDECEA", "FFF3D6", "EEF0F4", "F4F6F9", "FFFFFF"
FONT, HEAD_FONT = "Calibri", "Cambria"

RR, OVAL, DIAMOND = MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.OVAL, MSO_SHAPE.DIAMOND
LEFT, CENTER, RIGHT = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
TOP, MIDDLE = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def add_runs(par, runs, size, color, bold, font):
    for run_spec in runs:
        text, opts = (run_spec, {}) if isinstance(run_spec, str) else run_spec
        run = par.add_run()
        run.text = text
        run.font.name = opts.get("font", font)
        run.font.size = Pt(opts.get("size", size))
        run.font.bold = opts.get("bold", bold)
        run.font.color.rgb = rgb(opts.get("color", color))


def fill_text(tf, paras, size=12, color=INK, bold=False, align=LEFT, anchor=TOP,
              margins=(0, 0, 0, 0), space_after=0, font=FONT):
    """paras: list of paragraphs; a paragraph is a string or a list of runs (text or (text, options))."""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Inches(m) for m in margins)
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        if space_after:
            par.space_after = Pt(space_after)
        add_runs(par, para if isinstance(para, list) else [para], size, color, bold, font)


def textbox(slide, x, y, w, h, paras, name, **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    fill_text(box.text_frame, paras, **kw)
    return box


def shape(slide, kind, x, y, w, h, fill, line, name, line_w=1.0, radius=None, paras=None, **kw):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None and kind == RR:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if paras is not None:
        fill_text(s.text_frame, paras, **kw)
    return s


def arrow(slide, x1, y1, x2, y2, name, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.name = name
    c.line.color.rgb = rgb(ARROW)
    c.line.width = Pt(1.75)
    if head:
        tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


def label(slide, x, y, w, text, name, color=MUTED, bold=False):
    return textbox(slide, x, y, w, 0.22, [text], name, size=10, color=color, bold=bold, align=CENTER, anchor=MIDDLE)


def lane(slide, x, w, color, tint, title, port):
    shape(slide, RR, x, 0.95, w, 6.05, tint, None, f"Lane: {title}", radius=0.15)
    textbox(slide, x + 0.2, 1.0, w - 0.4, 0.35,
            [[(title.upper(), {"bold": True, "color": color}), ("    " + port, {"size": 11, "color": MUTED})]],
            f"Lane title: {title}", size=13, anchor=MIDDLE)


def card(slide, x, y, w, h, color, name):
    return shape(slide, RR, x, y, w, h, WHITE, color, name, line_w=1.25, radius=0.1)


def badge(slide, x, y, n, color):
    return shape(slide, OVAL, x, y, 0.3, 0.3, color, None, f"Step {n}", paras=[str(n)],
                 size=12, bold=True, color=WHITE, align=CENTER, anchor=MIDDLE)


def title_of(slide, x, y, w, text, name):
    return textbox(slide, x, y, w, 0.3, [text], name, size=14, bold=True, anchor=MIDDLE)


def chip(slide, x, y, w, h, text, fill, line, color, name, bold=True):
    return shape(slide, RR, x, y, w, h, fill, line, name, radius=0.06, paras=[text],
                 size=10, bold=bold, color=color, align=CENTER, anchor=MIDDLE)


def mock_row(slide, x, y, w, text, tag, name):
    shape(slide, RR, x, y, w, 0.26, ROW, None, name, radius=0.05, paras=[text], size=10,
          anchor=MIDDLE, margins=(0.08, 0, 0.08, 0))
    return tag


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = "Reimbursement System: Block Diagram"
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only

    # ---- Title -----------------------------------------------------------------
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = Inches(0.5), Inches(0.2), Inches(12.33), Inches(0.65)
    fill_text(title.text_frame, ["Reimbursement System: Block Diagram"], size=32, bold=True,
              anchor=MIDDLE, font=HEAD_FONT)

    # ---- Lanes -----------------------------------------------------------------
    lane(slide, 0.5, 3.2, BLUE, BLUE_T, "Employee portal", ":5001")
    lane(slide, 4.25, 4.7, GREEN, GREEN_T, "Shared backend", "Flask API + SQLite  :5000")
    lane(slide, 9.5, 3.33, PURPLE, PURPLE_T, "HR portal", ":5002")

    # ---- Employee lane ---------------------------------------------------------
    card(slide, 0.65, 1.45, 2.9, 1.0, BLUE, "Card: Upload bills")
    badge(slide, 0.77, 1.55, 1, BLUE)
    title_of(slide, 1.17, 1.55, 2.3, "Upload bills", "Title: Upload bills")
    textbox(slide, 0.8, 1.9, 2.6, 0.5, ["Photos or PDFs, many at once. Each bill becomes one row."],
            "Text: Upload bills")

    card(slide, 0.65, 2.65, 2.9, 2.1, BLUE, "Card: Review and send")
    badge(slide, 0.77, 2.75, 3, BLUE)
    title_of(slide, 1.17, 2.75, 2.3, "Review & send", "Title: Review and send")
    textbox(slide, 0.8, 3.1, 2.6, 0.62, ["Fix any field the reader got wrong, then send each bill to HR."],
            "Text: Review and send")
    for i, text in enumerate(["Bill 1 · ₹1,980", "Bill 2 · ₹2,050", "Bill 3 · ₹9,450"]):
        y = 3.8 + i * 0.3
        mock_row(slide, 0.8, y, 2.6, text, None, f"Mock row {i + 1}")
        chip(slide, 2.78, y + 0.03, 0.55, 0.2, "Send", BLUE, None, WHITE, f"Mock Send button {i + 1}")

    card(slide, 0.65, 5.2, 2.9, 1.65, BLUE, "Card: Track status")
    badge(slide, 0.77, 5.3, 9, BLUE)
    title_of(slide, 1.17, 5.3, 2.3, "Track status", "Title: Track status")
    textbox(slide, 0.8, 5.65, 2.6, 0.6, ["HR's decision shows up in seconds, with the reason if rejected."],
            "Text: Track status")
    for i, (text, fill, color) in enumerate([("Pending", AMBER_T, AMBER), ("Approved", GREEN_T, GREEN),
                                             ("Rejected", RED_T, RED)]):
        chip(slide, 0.8 + i * 0.92, 6.33, 0.82, 0.3, text, fill, None, color, f"Status chip {text}")

    # ---- Backend lane ----------------------------------------------------------
    card(slide, 4.4, 1.45, 4.4, 1.35, GREEN, "Card: Bill reader")
    badge(slide, 4.52, 1.55, 2, GREEN)
    title_of(slide, 5.04, 1.55, 2.3, "Bill reader", "Title: Bill reader")
    chip(slide, 7.55, 1.57, 1.1, 0.26, "AI / OCR", GREEN, None, WHITE, "Chip: AI / OCR")
    textbox(slide, 4.55, 1.9, 4.1, 0.5,
            ["Claude AI if an API key is set, else Tesseract OCR. Also tags the bill type."],
            "Text: Bill reader")
    for i, text in enumerate(["Bill no.", "Date", "Amount", "Type"]):
        chip(slide, 4.55 + i * 1.0, 2.43, 0.9, 0.28, text, GREEN_T, None, GREEN, f"Parsed field: {text}")

    card(slide, 4.4, 3.05, 4.4, 0.4, GREEN, "Card: Checks on send")
    badge(slide, 4.47, 3.1, 4, GREEN)
    textbox(slide, 5.0, 3.05, 3.7, 0.4,
            [[("Checks on send: ", {"bold": True}), "fields, date, duplicates"]],
            "Text: Checks on send", anchor=MIDDLE)

    card(slide, 4.4, 3.7, 4.4, 1.25, RED, "Card: Anomaly detection")
    badge(slide, 4.52, 3.8, 5, RED)
    title_of(slide, 5.04, 3.8, 2.5, "Anomaly detection", "Title: Anomaly detection")
    chip(slide, 7.65, 3.82, 1.0, 0.26, "ML model", RED, None, WHITE, "Chip: ML model")
    textbox(slide, 4.55, 4.15, 4.1, 0.75,
            [[("Rules: ", {"bold": True}), "amount vs own history, company-wide and monthly total"],
             [("ML: ", {"bold": True}), "Isolation Forest learns from approved bills"]],
            "Text: Anomaly detection", space_after=3)

    shape(slide, DIAMOND, 4.9, 5.2, 3.4, 0.9, AMBER_T, AMBER, "Decision: auto-approve", line_w=1.25,
          paras=["Auto-approve ON and no anomaly?"], size=11, bold=True, align=CENTER, anchor=MIDDLE)
    badge(slide, 4.5, 5.5, 6, AMBER)

    shape(slide, RR, 4.4, 6.4, 4.4, 0.5, WHITE, GREY, "Shared database", line_w=1.25, radius=0.1,
          paras=[[("SQLite database  ", {"bold": True}), "· bills, users, decisions, settings"]],
          size=11, align=CENTER, anchor=MIDDLE)

    # ---- HR lane ---------------------------------------------------------------
    card(slide, 9.65, 1.45, 3.03, 2.5, PURPLE, "Card: Waiting for review")
    badge(slide, 9.77, 1.55, 7, PURPLE)
    title_of(slide, 10.27, 1.55, 2.3, "Waiting for review", "Title: Waiting for review")
    textbox(slide, 9.8, 1.9, 2.73, 0.62, ["Flagged bills show red with reasons. Rejecting needs a reason."],
            "Text: Waiting for review")
    shape(slide, RR, 9.8, 2.6, 0.42, 0.22, GREEN, None, "Mock switch track", radius=0.11)
    shape(slide, OVAL, 10.02, 2.62, 0.18, 0.18, WHITE, None, "Mock switch knob")
    textbox(slide, 10.3, 2.58, 2.2, 0.26, ["Auto-approve switch: ON"], "Mock switch label", size=10, anchor=MIDDLE)
    shape(slide, RR, 9.8, 2.95, 2.73, 0.5, RED_T, RED, "Mock anomaly row", radius=0.06)
    textbox(slide, 9.88, 2.97, 1.55, 0.46,
            [[("Amit · ₹9,450", {"bold": True, "color": RED, "size": 11})],
             [("2.8× usual Fuel bill", {"color": RED})]],
            "Mock anomaly text", size=10, anchor=MIDDLE)
    chip(slide, 11.5, 3.08, 0.95, 0.24, "⚠ Anomaly", RED, None, WHITE, "Mock anomaly tag")
    chip(slide, 9.8, 3.53, 0.85, 0.26, "Approve", GREEN, None, WHITE, "Mock Approve button")
    chip(slide, 10.73, 3.53, 0.75, 0.26, "Reject", WHITE, RED, RED, "Mock Reject button")

    card(slide, 9.65, 4.4, 3.03, 2.45, PURPLE, "Card: Decision history")
    badge(slide, 9.77, 4.5, 8, PURPLE)
    title_of(slide, 10.27, 4.5, 2.3, "Decision history", "Title: Decision history")
    textbox(slide, 9.8, 4.85, 2.73, 0.62,
            ["Every approved and rejected bill stays listed, with who decided and when."],
            "Text: Decision history")
    for i, (text, tag) in enumerate([("₹2,100 · Approved", "Auto-approved"), ("₹9,450 · Rejected", "Manual"),
                                     ("₹1,860 · Approved", "Manual")]):
        y = 5.6 + i * 0.3
        mock_row(slide, 9.8, y, 2.73, text, tag, f"History row {i + 1}")
        if tag == "Manual":
            chip(slide, 11.83, y + 0.03, 0.62, 0.2, tag, WHITE, GREY, MUTED, f"History tag {i + 1}")
        else:
            chip(slide, 11.47, y + 0.03, 0.98, 0.2, tag, BLUE_T, BLUE, BLUE, f"History tag {i + 1}")

    # ---- Arrows ----------------------------------------------------------------
    arrow(slide, 3.55, 1.9, 4.4, 1.9, "Arrow: upload")
    label(slide, 3.55, 1.64, 0.85, "Upload", "Label: upload")
    arrow(slide, 4.4, 2.7, 3.55, 2.7, "Arrow: parsed bill")
    label(slide, 3.55, 2.44, 0.85, "Parsed bill", "Label: parsed bill")
    arrow(slide, 3.55, 3.25, 4.4, 3.25, "Arrow: send to HR")
    label(slide, 3.55, 2.99, 0.85, "Send to HR", "Label: send to HR")
    arrow(slide, 6.6, 3.45, 6.6, 3.7, "Arrow: checks to anomaly")
    arrow(slide, 6.6, 4.95, 6.6, 5.2, "Arrow: anomaly to decision")
    arrow(slide, 8.3, 5.65, 9.225, 5.65, "Line: decision out", head=False)
    arrow(slide, 9.225, 3.4, 9.225, 6.1, "Line: decision split", head=False)
    arrow(slide, 9.225, 3.4, 9.65, 3.4, "Arrow: to HR review")
    label(slide, 9.27, 3.14, 0.35, "No", "Label: no", bold=True)
    arrow(slide, 9.225, 6.1, 9.65, 6.1, "Arrow: auto-approved to history")
    label(slide, 9.27, 5.84, 0.35, "Yes", "Label: yes", color=GREEN, bold=True)
    arrow(slide, 11.165, 3.95, 11.165, 4.4, "Arrow: review to history")
    arrow(slide, 6.6, 6.1, 6.6, 6.4, "Arrow: decision to database")
    arrow(slide, 9.65, 6.65, 8.8, 6.65, "Arrow: decisions to database")
    label(slide, 8.8, 6.4, 0.85, "Decision", "Label: decision")
    arrow(slide, 4.4, 6.65, 3.55, 6.65, "Arrow: status to employee")
    label(slide, 3.55, 6.4, 0.85, "Status", "Label: status")

    slide.notes_slide.notes_text_frame.text = (
        "Walk-through: (1) The employee uploads photos or PDFs of fuel, repair and maintenance bills. "
        "(2) The bill reader fills in bill number, date, amount and bill type: Claude AI if an API key is set, "
        "otherwise local Tesseract OCR. (3) The parsed fields appear as one row per bill; the employee fixes anything "
        "wrong and sends the bill to HR. (4) The backend checks required fields, a valid date and duplicates. "
        "(5) Anomaly detection compares the amount with history using statistical rules and an Isolation Forest model, "
        "both learned only from this system's own approved bills. (6) If HR's auto-approve switch is on and the bill is "
        "not flagged, it is approved automatically; otherwise it goes to HR. (7) HR sees flagged bills in red with the "
        "reasons, then approves or rejects (a reason is required to reject). (8) Every decision stays in the history, "
        "tagged auto-approved or manual. (9) The employee sees Pending, Approved or Rejected within seconds. "
        "Both portals talk to the same backend and database.")

    prs.save(OUT)
    print("Saved", OUT)


def export_preview():
    script = (
        "$ErrorActionPreference='Stop';"
        f"$app=New-Object -ComObject PowerPoint.Application;"
        f"$p=$app.Presentations.Open('{OUT}',$true,$false,$false);"
        f"$p.Slides.Item(1).Export('{PREVIEW}','PNG',1920,1080);"
        "$p.Close();$app.Quit()"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)
    print("Preview saved", PREVIEW)


if __name__ == "__main__":
    build()
    if "--preview" in sys.argv:
        export_preview()
